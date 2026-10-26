"""Regression tests for file imports, optional identifiers and Word round trips."""
from __future__ import annotations
import shutil
import tempfile
import unittest
import uuid
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from schoolregister.excel import (parse_file, export_excel, asset_template,
                                  asset_tajik_template, asset_tajik_no_number_template, save_template)
from schoolregister.word import export_word, save_word_template, _office_path
from schoolregister.storage import Database
from schoolregister.constants import HEADERS, TAJIK_HEADERS


class DocumentInterchangeV3(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path=Path(self.tmp.name)
        self.db=Database(self.path/'school.db')
        self.addCleanup(self.db.close)

    def write_xlsx(self, filename, headers, rows):
        wb=Workbook();ws=wb.active
        ws.append(headers)
        for row in rows:ws.append(row)
        name=self.path/filename
        wb.save(name)
        return name

    def write_docx(self, filename, headers, rows):
        doc=Document();doc.add_paragraph('Школьный журнал')
        table=doc.add_table(rows=1,cols=len(headers))
        for index,h in enumerate(headers):table.rows[0].cells[index].text=str(h)
        for row in rows:
            cells=table.add_row().cells
            for index,item in enumerate(row):cells[index].text=str(item or '')
        name=self.path/filename
        doc.save(name)
        return name

    def check_unique_ids(self):
        rows=self.db.con.execute('SELECT id FROM students').fetchall()
        self.assertEqual(len(rows),len(set(r['id'] for r in rows)))
        for row in rows:
            self.assertEqual(str(uuid.UUID(row['id'])),row['id'])

    def test_tajik_xlsx_5_columns_without_number_generates_id_and_history(self):
        p=self.write_xlsx('school_without_num.xlsx',TAJIK_HEADERS[1:],[
            ['Раҳимов Али','5А','+992901234567','Мактаби 10','Мактаби 4'],
            ['Каримова Ситора','6Б','+992991234567','','Мактаби 7'],
        ])
        result=parse_file(p)
        self.assertEqual(len(result.records),2)
        self.assertTrue(all(not r['id'] for r in result.records))
        self.assertEqual(self.db.apply_import(result.records),(2,0))
        self.assertEqual(self.db.apply_import(result.records),(0,0))
        self.check_unique_ids()
        rows=self.db.con.execute('SELECT * FROM students ORDER BY grade').fetchall()
        self.assertEqual(rows[0]['status'],'Переведён')
        self.assertEqual(rows[0]['to_school'],'Мактаби 10')
        self.assertEqual(len(self.db.transfer_history(rows[0]['id'])),2)

    def test_tajik_xlsx_6_columns_ignores_number_and_shared_across_formats(self):
        row=['Саидов Фирдавс','9','+992920000001','', 'Мактаби 2']
        a=self.write_xlsx('with_num.xlsx',TAJIK_HEADERS,[[999,*row]])
        b=self.write_docx('without_num.docx',TAJIK_HEADERS[1:],[row])
        ra=parse_file(a);rb=parse_file(b)
        self.assertEqual(ra.records[0]['_fingerprint'],rb.records[0]['_fingerprint'])
        self.assertEqual(self.db.apply_import(ra.records),(1,0))
        self.assertEqual(self.db.apply_import(rb.records),(0,0))
        self.check_unique_ids()

    def test_full_xlsx_and_word_without_id(self):
        record=['Раҳимов','Али','', '2014-01-01',5,'А','', 'Душанбе',
                'Раҳимова','+992901111111','','','Учится','','','','']
        self.assertEqual(len(record),17)
        for ext in ('xlsx','docx'):
            if ext=='xlsx':
                path=self.write_xlsx('full.xlsx',HEADERS[1:],[record])
            else:
                path=self.write_docx('full.docx',HEADERS[1:],[record])
            result=parse_file(path)
            self.assertEqual(result.template,'modern')
            self.assertEqual(len(result.records),1)
            self.assertEqual(result.records[0]['id'],'')
            self.assertEqual(self.db.apply_import(result.records),(1,0) if ext=='xlsx' else (0,0))
        self.check_unique_ids()

    def test_full_duplicate_shared_phone_without_dob_distinct_students(self):
        records=[]
        for class_number in ('5','6'):
            vals=['Раҳимов','Али','','',class_number,'А','','', 'Раҳимов',
                  '+992901111111','','','Учится','','','','']
            records.append(vals)
        r=parse_file(self.write_xlsx('names.xlsx', HEADERS[1:],records))
        self.assertEqual(len(r.records),2)
        self.assertEqual(self.db.apply_import(r.records),(2,0))
        self.assertEqual(self.db.list_students()[1],2)
        self.check_unique_ids()

    def test_all_packaged_templates_are_optional_id(self):
        for path in (asset_template(),asset_tajik_template(),asset_tajik_no_number_template()):
            preview=parse_file(path)
            self.assertEqual(preview.errors,[])
            self.assertEqual(len(preview.records),0)
        for ext in ('xlsx','docx'):
            target=self.path/('template.'+ext)
            save_template(target)
            p=parse_file(target)
            self.assertEqual(p.errors,[])
            self.assertEqual(p.records,[])

    def test_word_export_and_import_all_variants(self):
        self.db.apply_import(parse_file(self.write_xlsx('basic.xlsx',TAJIK_HEADERS[1:],[
            ['Назаров Давлат','7А','+992901234567','Мактаби 1','Мактаби 12']] )).records)
        for tajik in (False,True):
            for no_number in ((False,True) if tajik else (False,)):
                target=self.path/(f'export_{tajik}_{no_number}.docx')
                export_word(self.db,target,tajik=tajik,no_number=no_number)
                result=parse_file(target)
                self.assertEqual(len(result.records),1)
                self.assertEqual(result.records[0]['from_school'],'Мактаби 12')
                self.assertEqual(result.records[0]['to_school'],'Мактаби 1')
                if not tajik:self.assertTrue(result.records[0]['id'])
                fresh=Database(self.path/f'fresh_{tajik}_{no_number}.db')
                try:
                    self.assertEqual(fresh.apply_import(result.records),(1,0))
                    self.assertEqual(fresh.list_students()[1],1)
                finally:fresh.close()

    @unittest.skipUnless(shutil.which('antiword') and _office_path(), 'DOC needs installed converter + parser for test')
    def test_legacy_doc_actual_format_import_export(self):
        save_word_template(self.path/'empty.doc',tajik=True,no_number=True)
        self.assertEqual(parse_file(self.path/'empty.doc').records,[])
        p=self.write_docx('filled.docx',TAJIK_HEADERS[1:],[
            ['Назаров Давлат','5А','+992901234567','Мактаби 4','Мактаби 1']])
        from schoolregister.word import convert_legacy_word
        doc=self.path/'filled.doc'
        convert_legacy_word(p,doc)
        result=parse_file(doc)
        self.assertEqual(len(result.records),1)
        self.assertEqual(self.db.apply_import(result.records),(1,0))
        path=self.path/'out.doc'
        export_word(self.db,path,tajik=True,no_number=True)
        self.assertEqual(len(parse_file(path).records),1)

    def test_invalid_word_no_table_and_formula_are_rejected(self):
        doc=Document();doc.add_paragraph('Раҳимов Али - 5А')
        p=self.path/'plain.docx';doc.save(p)
        with self.assertRaisesRegex(ValueError,'нет таблиц'):
            parse_file(p)
        p=self.write_xlsx('formula.xlsx',TAJIK_HEADERS[1:],[[
            'Раҳимов Али','=2+3','+992901234567','','']])
        self.assertIn('формулы запрещены',parse_file(p).errors[0])


if __name__=='__main__':unittest.main()
