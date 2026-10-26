"""Regression tests for Tajik school sheets, transfers, annual promotion and migration."""
from __future__ import annotations
import sqlite3
import tempfile
import unittest
from datetime import date
from pathlib import Path
from openpyxl import Workbook, load_workbook
from docx import Document

from schoolregister.excel import parse_excel, export_excel, asset_template, asset_tajik_template
from schoolregister.storage import Database, school_year
from schoolregister.validation import clean_record


def pupil(**extras):
    raw={'surname':'Раҳимов','given_name':'Али','date_of_birth':'2016-01-03',
         'grade':5,'address':'Ноҳияи Сино','guardian_phone':'+992900001234',
         'status':'Учится'}
    raw.update(extras)
    return clean_record(raw,allow_incomplete=True)


class SchoolV2Tests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.folder=Path(self.tmp.name)
        self.db=Database(self.folder/'school.db');self.addCleanup(self.db.close)

    def test_new_and_old_templates(self):
        self.assertEqual(parse_excel(asset_template()).template,'modern')
        self.assertEqual(parse_excel(asset_tajik_template()).template,'tajik')
        self.assertEqual(parse_excel(asset_tajik_template()).records,[])

    def tajik_file(self, *, as_word=False, rows=None):
        heading=['№','Ном ва насаб','Синф','Рақами телефон',
                 'Рафт (ба кадом мактаб)','Омад (аз кадом мактаб)']
        rows=rows or [[1,'Раҳимов Али','5А','+992900001234','Мактаби № 9','Мактаби № 3']]
        if as_word:
            doc=Document();table=doc.add_table(rows=1,cols=6)
            for i,s in enumerate(heading):table.rows[0].cells[i].text=s
            for row in rows:
                cells=table.add_row().cells
                for i,s in enumerate(row):cells[i].text=str(s)
            path=self.folder/'school.docx';doc.save(path)
        else:
            wb=Workbook();ws=wb.active;ws.title='Ҳисобот';ws.append(['Рӯйхати хонандагон'])
            ws.append(heading)
            for row in rows:ws.append(row)
            path=self.folder/'school.xlsx';wb.save(path)
        return path

    def test_tajik_xlsx_transfer_fields_history_and_stats(self):
        data=parse_excel(self.tajik_file())
        self.assertFalse(data.errors)
        self.assertEqual(data.template,'tajik')
        self.assertEqual(len(data.records),1)
        self.assertEqual(data.records[0]['section'],'А')
        self.assertEqual(data.records[0]['date_of_birth'],'')
        self.assertEqual(data.records[0]['status'],'Переведён')
        self.assertEqual(self.db.apply_import(data.records),(1,0))
        self.assertEqual(self.db.apply_import(data.records),(0,0))
        s=self.db.list_students()[0][0]
        self.assertEqual(s['from_school'],'Мактаби № 3')
        self.assertEqual(s['to_school'],'Мактаби № 9')
        history=self.db.transfer_history(s['id'])
        self.assertEqual({e['direction'] for e in history},{'in','out'})
        self.assertTrue(all(e['recorded_at'] and not e['event_date'] for e in history))
        self.assertEqual(self.db.stats()['studying'],0)
        self.assertEqual(self.db.stats()['transferred'],1)
        self.assertEqual(self.db.promote_for_school_year(date(2030,9,4)),0)
        self.assertEqual(self.db.get_student(s['id'])['grade'],5)

    def test_tajik_word_table_import(self):
        data=parse_excel(self.tajik_file(as_word=True))
        self.assertEqual(len(data.records),1)
        self.assertEqual(self.db.apply_import(data.records),(1,0))

    def test_tajik_title_blank_rows_and_grade_text(self):
        data=parse_excel(self.tajik_file(rows=[['','','','','',''],
            [2,'Саидова Мадина','10-Б','991112233','','Мактаби № 4']]))
        self.assertEqual(len(data.records),1)
        self.assertEqual(data.records[0]['grade'],10)
        self.assertEqual(data.records[0]['section'],'Б')
        self.assertEqual(data.records[0]['status'],'Учится')
        self.assertEqual(data.records[0]['from_school'],'Мактаби № 4')

    def test_tajik_localized_address_headers_import_to_database(self):
        headings=['Рақами тартиби','Суроғаи истиқомат','Телефони волидайн','Синф','Ном ва насаб']
        values=[1,'Душанбе, ноҳияи Сино','+992900001234','5-Ҳ','Раҳимов Али']
        wb=Workbook();ws=wb.active;ws.append(headings);ws.append(values)
        path=self.folder/'tajik_address.xlsx';wb.save(path)

        preview=parse_excel(path)
        self.assertFalse(preview.errors)
        self.assertEqual(len(preview.records),1)
        self.assertEqual(preview.records[0]['address'],'Душанбе, ноҳияи Сино')
        self.assertEqual(preview.records[0]['section'],'Ҳ')
        self.assertEqual(self.db.apply_import(preview.records),(1,0))
        self.assertEqual(self.db.list_students()[0][0]['address'],'Душанбе, ноҳияи Сино')

        doc=Document();table=doc.add_table(rows=1,cols=len(headings))
        for col,heading in enumerate(headings):table.rows[0].cells[col].text=heading
        cells=table.add_row().cells
        for col,value in enumerate(values):cells[col].text=str(value)
        word_path=self.folder/'tajik_address.docx';doc.save(word_path)
        word_preview=parse_excel(word_path)
        self.assertFalse(word_preview.errors)
        self.assertEqual(word_preview.records[0]['address'],'Душанбе, ноҳияи Сино')

    def test_word_imports_all_tables_with_unaccented_tajik_address_heading(self):
        doc=Document()
        heading=['№','Ном ва насаб','Синф','Чои истикомат','Рақами телефон']
        entries=[
            ['1','Раҳимов Али','5А','Душанбе, ноҳияи Сино','+992900000001'],
            ['2','Каримова Мадина','6Б','Ҳисор, марказ','+992900000002'],
        ]
        for entry in entries:
            table=doc.add_table(rows=1,cols=len(heading))
            for col,value in enumerate(heading):table.rows[0].cells[col].text=value
            cells=table.add_row().cells
            for col,value in enumerate(entry):cells[col].text=value
        path=self.folder/'realistic_tajik_address.docx';doc.save(path)

        preview=parse_excel(path)
        self.assertFalse(preview.errors)
        self.assertEqual(len(preview.records),2)
        self.assertEqual(
            {record['address'] for record in preview.records},
            {'Душанбе, ноҳияи Сино','Ҳисор, марказ'},
        )
        self.assertEqual(self.db.apply_import(preview.records),(2,0))
        self.assertEqual(
            {row['address'] for row in self.db.list_students()[0]},
            {'Душанбе, ноҳияи Сино','Ҳисор, марказ'},
        )

    def test_tajik_word_import_keeps_address_when_phone_has_extra_letter(self):
        doc=Document()
        table=doc.add_table(rows=1,cols=5)
        heading=['№','Ном ва насаб','Синф','Чои истикомат','Рақами телефон']
        for col,value in enumerate(heading):table.rows[0].cells[col].text=value
        row=['1','Зуҳурова Таиба Саидмуминовна','6Б','К. Соҳилӣ 5','900-01-51-62ь']
        cells=table.add_row().cells
        for col,value in enumerate(row):cells[col].text=value
        path=self.folder/'phone_typo_tajik.docx';doc.save(path)

        preview=parse_excel(path)
        self.assertFalse(preview.errors)
        self.assertEqual(len(preview.records),1)
        self.assertEqual(preview.records[0]['address'],'К. Соҳилӣ 5')
        self.assertEqual(preview.records[0]['guardian_phone'],'')
        self.assertTrue(any('Строка 2' in warning for warning in preview.warnings))
        self.assertEqual(self.db.apply_import(preview.records),(1,0))
        self.assertEqual(self.db.list_students()[0][0]['address'],'К. Соҳилӣ 5')

    def test_formulas_in_import_rejected_except_pre_numbered_tajik_template(self):
        wb=load_workbook(asset_tajik_template());ws=wb.worksheets[0]
        ws['B2']='Раҳимов Али';ws['C2']='5';ws['D2']='+992900001234'
        target=self.folder/'filled.xlsx';wb.save(target)
        preview=parse_excel(target)
        self.assertEqual(len(preview.records),1)
        self.assertFalse(preview.errors)
        ws['C2']='=2+3';wb.save(target)
        preview=parse_excel(target)
        self.assertEqual(len(preview.records),0)
        self.assertIn('формулы запрещены',preview.errors[0])

    def test_promotion_once_per_year_missed_years_and_manual_class_edits(self):
        rid=self.db.save_student(pupil())
        self.db.con.execute('UPDATE students SET grade_year=2023 WHERE id=?',(rid,))
        self.db.con.commit()
        self.assertEqual(self.db.promote_for_school_year(date(2026,8,31)),1)
        self.assertEqual(self.db.get_student(rid)['grade'],7)
        self.assertEqual(self.db.promote_for_school_year(date(2026,8,31)),0)
        self.assertEqual(self.db.promote_for_school_year(date(2026,9,1)),1)
        self.assertEqual(self.db.get_student(rid)['grade'],8)
        self.assertEqual(self.db.promote_for_school_year(date(2026,9,1)),0)
        self.assertEqual(len(self.db.grade_history(rid)),2)
        self.db.save_student(pupil(grade=9),rid)
        self.assertEqual(self.db.get_student(rid)['grade_year'],school_year())
        self.assertEqual(self.db.promote_for_school_year(date(2026,10,1)),0)

    def test_terminal_grade_review_not_autograduated(self):
        rid=self.db.save_student(pupil(grade=11))
        self.db.con.execute('UPDATE students SET grade_year=2025 WHERE id=?',(rid,));self.db.con.commit()
        self.assertEqual(self.db.promote_for_school_year(date(2026,9,1)),0)
        self.assertEqual(self.db.stats()['review'],1)
        self.assertEqual(self.db.get_student(rid)['status'],'Учится')
        self.db.set_final_grade(12)
        self.assertEqual(self.db.promote_for_school_year(date(2026,9,1)),1)
        self.assertEqual(self.db.get_student(rid)['grade'],12)

    def test_history_is_append_only_and_count_by_grade(self):
        rid=self.db.save_student(pupil(from_school='Школа 1'))
        self.db.save_student(pupil(from_school='Школа 2',to_school='Школа 3',left_on='2026-09-12'),rid)
        history=self.db.transfer_history(rid)
        self.assertEqual(len(history),3)
        self.assertEqual(history[0]['direction'],'out')
        self.assertEqual(history[0]['event_date'],'2026-09-12')
        self.assertEqual(self.db.stats()['arriving'],1)
        self.assertFalse(self.db.grade_distribution())
        self.assertTrue(self.db.soft_delete(rid))
        self.assertTrue(self.db.restore(rid))
        self.assertEqual(len(self.db.transfer_history(rid)),3)

    def test_export_roundtrip_in_both_formats(self):
        self.db.save_student(pupil(from_school='Школа 7',to_school='Школа 9'))
        for tajik in (False,True):
            path=self.folder/f'export_{tajik}.xlsx'
            export_excel(self.db,path,tajik=tajik)
            preview=parse_excel(path)
            self.assertFalse(preview.errors,preview.errors)
            self.assertEqual(len(preview.records),1)
            self.assertEqual(preview.records[0]['to_school'],'Школа 9')

    def test_duplicate_names_without_dob_and_phone_not_dropped(self):
        records=[pupil(date_of_birth='',address='',guardian_phone='',surname='Набиев'),
                 pupil(date_of_birth='',address='',guardian_phone='',surname='Набиев')]
        self.assertEqual(self.db.apply_import(records),(2,0))

    def test_v1_db_migration_retains_student(self):
        other=self.folder/'from_v1.db'
        con=sqlite3.connect(other)
        con.executescript('''CREATE TABLE students(
           id TEXT PRIMARY KEY,surname TEXT NOT NULL,given_name TEXT NOT NULL,patronymic TEXT NOT NULL DEFAULT '',
           date_of_birth TEXT NOT NULL,grade INTEGER NOT NULL,section TEXT NOT NULL DEFAULT '',shift TEXT NOT NULL DEFAULT '',
           address TEXT NOT NULL,guardian_name TEXT NOT NULL DEFAULT '',guardian_phone TEXT NOT NULL DEFAULT '',
           extra_phone TEXT NOT NULL DEFAULT '',admitted_on TEXT NOT NULL DEFAULT '',status TEXT NOT NULL DEFAULT 'Учится',
           notes TEXT NOT NULL DEFAULT '',created_at TEXT NOT NULL,updated_at TEXT NOT NULL,deleted_at TEXT);
           CREATE TABLE settings(key TEXT PRIMARY KEY,value TEXT NOT NULL);
           CREATE TABLE audit_log(id INTEGER PRIMARY KEY AUTOINCREMENT,at TEXT NOT NULL,action TEXT NOT NULL,
                 record_id TEXT,detail TEXT NOT NULL DEFAULT '');
           PRAGMA user_version=1;''')
        con.execute('''INSERT INTO students(id,surname,given_name,date_of_birth,grade,address,created_at,updated_at)
          VALUES('from-old','Иванов','Иван','2015-01-01',5,'Душанбе','2026-01-01','2026-01-01')''')
        con.commit();con.close()
        migrated=Database(other)
        try:
            self.assertEqual(migrated.con.execute('PRAGMA user_version').fetchone()[0],2)
            self.assertEqual(migrated.get_student('from-old')['grade'],5)
            self.assertEqual(migrated.get_student('from-old')['grade_year'],school_year())
            self.assertEqual(migrated.promote_for_school_year(date.today()),0)
        finally:migrated.close()

if __name__=='__main__':unittest.main()

class AdditionalV2SafetyTests(unittest.TestCase):
    def test_older_excel_import_does_not_erase_new_transfer_fields(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/'data.db')
            try:
                rid=db.save_student(pupil(from_school='School A'))
                wb=Workbook();ws=wb.active
                from schoolregister.constants import HEADERS
                ws.append(HEADERS[:15]);ws.append([rid,'Раҳимов','Али','','2016-01-03',5,'А','',
                                                    'Новый адрес','','+992900001234','','','Учится',''])
                path=Path(folder)/'old.xlsx';wb.save(path)
                prev=parse_excel(path)
                self.assertFalse(prev.errors)
                self.assertEqual(db.apply_import(prev.records,'update'),(0,1))
                self.assertEqual(db.get_student(rid)['from_school'],'School A')
            finally:db.close()

    def test_transfer_date_revisions_preserve_previous_event(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Database(Path(folder)/'data.db')
            try:
                rid=db.save_student(pupil(to_school='School B'))
                before=db.transfer_history(rid)
                db.save_student(pupil(to_school='School B',left_on='2026-09-01'),rid)
                after=db.transfer_history(rid)
                self.assertEqual(len(after),len(before)+1)
                self.assertEqual(after[0]['event_date'],'2026-09-01')
                self.assertEqual(after[1]['event_date'],'')
                self.assertEqual(db.list_students(status='Переведён')[1],1)
                self.assertEqual(db.list_students(status='Учится')[1],0)
            finally:db.close()
