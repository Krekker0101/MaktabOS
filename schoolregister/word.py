"""Document interchange for real Word tables. DOC is handled by local Word/LibreOffice."""
from __future__ import annotations
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import TYPE_CHECKING
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor

from .constants import FIELDS, HEADERS, MAX_IMPORT_ROWS, TAJIK_HEADERS
if TYPE_CHECKING:
    from .excel import ImportPreview


def _office_path() -> str | None:
    candidate = shutil.which('libreoffice') or shutil.which('soffice')
    if candidate:
        return candidate
    if sys.platform == 'win32':
        bases = [os.environ.get('PROGRAMFILES', ''), os.environ.get('PROGRAMFILES(X86)', '')]
        for root in filter(None,bases):
            for name in ('LibreOffice', 'LibreOfficePortable'):
                check=Path(root)/name/'program'/'soffice.exe'
                if check.exists():return str(check)
    return None


def _convert_libreoffice(source: Path, output: Path) -> bool:
    executable=_office_path()
    if not executable:return False
    with tempfile.TemporaryDirectory(prefix='schoolregister_lo_') as temp:
        folder=Path(temp)
        # unique office profile prevents conflicting running LO instances.
        profile=(folder/'profile').as_uri()
        fmt = 'docx:Office Open XML Text' if output.suffix.lower()=='.docx' else 'doc:MS Word 97'
        command=[executable,f'-env:UserInstallation={profile}','--headless','--convert-to',fmt,
                 '--outdir',str(folder),str(source.resolve())]
        env=os.environ.copy()
        env['SAL_USE_VCLPLUGIN']='svp'
        env.pop('DISPLAY', None)  # Headless conversion on Linux without X server
        try:
            proc=subprocess.run(command,check=False,timeout=75,stdout=subprocess.PIPE,stderr=subprocess.PIPE,env=env)
        except (OSError,subprocess.TimeoutExpired) as exc:
            raise ValueError('Не удалось преобразовать старый Word через LibreOffice (время ожидания/запуск).') from exc
        expected=folder/(source.stem+output.suffix.lower())
        if proc.returncode or not expected.is_file() or not expected.stat().st_size:
            details=proc.stderr.decode('utf-8','replace')[:250]
            raise ValueError(f'LibreOffice не смог преобразовать файл Word. {details}')
        shutil.copy2(expected,output)
        return True


def _convert_msword(source: Path, output: Path) -> bool:
    if sys.platform != 'win32':return False
    try:
        from win32com.client import DispatchEx
    except ImportError:
        return False
    app = None
    doc = None
    try:
        app=DispatchEx('Word.Application')
        app.Visible=False
        app.DisplayAlerts=0
        try: app.AutomationSecurity=3  # msoAutomationSecurityForceDisable
        except Exception:pass
        doc=app.Documents.Open(str(source.resolve()),ReadOnly=True,ConfirmConversions=False,
                               AddToRecentFiles=False,OpenAndRepair=True,NoEncodingDialog=True)
        doc.SaveAs2(str(output.resolve()),FileFormat=16 if output.suffix.lower()=='.docx' else 0,
                    AddToRecentFiles=False)
        return output.exists() and output.stat().st_size>0
    except Exception as exc:
        raise ValueError('Microsoft Word не смог преобразовать файл. Проверьте его целостность.') from exc
    finally:
        if doc is not None:
            try:doc.Close(SaveChanges=False)
            except Exception:pass
        if app is not None:
            try:app.Quit()
            except Exception:pass


def convert_legacy_word(source: Path, output: Path) -> None:
    """Legacy .doc needs an installed desktop converter; never silently rename file."""
    if source.suffix.lower()==output.suffix.lower():
        raise ValueError('Входной и выходной форматы одинаковы')
    if _convert_libreoffice(source,output):return
    if _convert_msword(source,output):return
    raise RuntimeError('Для Word .doc требуется установленный LibreOffice или Microsoft Word '
                       '(для Microsoft Word нужен дополнительный модуль pywin32). '
                       'Документы .docx работают без этих программ.')


def _parse_antiword(path: Path) -> 'ImportPreview | None':
    """Use Antiword's structured DocBook output; normal text extraction loses column boundaries."""
    executable=shutil.which('antiword')
    if not executable:
        return None
    from .excel import _rows_import
    try:
        run=subprocess.run([executable,'-x','db',str(path.resolve())],timeout=60,
                           stdout=subprocess.PIPE,stderr=subprocess.PIPE,check=False)
    except (OSError,subprocess.TimeoutExpired) as exc:
        raise ValueError('Старый DOC не удалось прочитать через Antiword') from exc
    if run.returncode or len(run.stdout)>50*1024*1024:
        raise ValueError('Невозможно прочитать Word DOC через Antiword или документ слишком большой')
    try:
        root=ET.fromstring(run.stdout)
    except ET.ParseError as exc:
        raise ValueError('Структура таблиц в DOC повреждена') from exc
    for table in root.findall('.//informaltable')+root.findall('.//table'):
        rows=[]
        for tr in table.findall('.//row'):
            rows.append([' '.join(''.join(x.itertext()).split()) for x in tr.findall('./entry')])
        try:
            return _rows_import(rows)
        except ValueError:
            pass
    raise ValueError('Не обнаружена таблица Word с известными заголовками')


def parse_word(path: Path) -> 'ImportPreview':
    from .excel import ImportPreview, _rows_import
    path=Path(path)
    if path.suffix.lower()=='.doc':
        antiword_result=_parse_antiword(path)
        if antiword_result is not None:
            return antiword_result
        with tempfile.TemporaryDirectory(prefix='schoolregister_doc_') as tmp:
            output=Path(tmp)/(path.stem+'.docx')
            # On Windows Microsoft Word preserves real tables; LibreOffice can
            # flatten particular DOC tables and cannot guarantee this.
            if sys.platform=='win32' and _convert_msword(path,output):
                return parse_word(output)
            convert_legacy_word(path,output)
            return parse_word(output)
    try:
        doc=Document(path)
    except Exception as exc:
        raise ValueError('Word-файл повреждён или не является документом .docx') from exc
    if not doc.tables:
        raise ValueError('В документе Word нет таблиц. Фото и сканы не поддерживаются.')
    previews = []
    total_rows = 0
    for table in doc.tables:
        if len(table.rows)>MAX_IMPORT_ROWS+31:
            raise ValueError('В Word больше 30 000 строк')
        rows=[[cell.text.strip() for cell in row.cells] for row in table.rows]
        try:
            preview = _rows_import(rows)
        except ValueError:
            continue
        total_rows += max(0, len(rows) - 1)
        if total_rows > MAX_IMPORT_ROWS:
            raise ValueError('В Word больше 30 000 строк')
        previews.append(preview)
    if not previews:
        raise ValueError('В Word нет таблицы с заголовками полного или таджикского шаблона')
    errors = [error for preview in previews for error in preview.errors]
    if len(errors) > 100:
        errors = errors[:100] + ['Слишком много ошибок (первые 100)']
    return ImportPreview(
        records=[record for preview in previews for record in preview.records],
        errors=errors,
        blank_rows=sum(preview.blank_rows for preview in previews),
        warnings=list(dict.fromkeys(
            warning for preview in previews for warning in preview.warnings
        )),
        template=previews[0].template,
    )


def _style_doc(doc: Document, tajik: bool) -> None:
    section=doc.sections[0]
    section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width=Cm(29.7 if tajik else 42)
    section.page_height=Cm(21 if tajik else 29.7)
    section.top_margin=Cm(1.4)
    section.bottom_margin=Cm(1.4)
    section.left_margin=Cm(1.3)
    section.right_margin=Cm(1.3)
    normal=doc.styles['Normal']
    normal.font.name='Calibri'
    normal.font.size=Pt(9 if tajik else 7)


def _table(doc, headers: list[str], rows: list[list[object]], tajik: bool):
    table=doc.add_table(rows=1, cols=len(headers))
    table.style='Table Grid'
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    widths = ([1.0,8,2,4.1,5.9,5.9] if len(headers)==6 and tajik else
              [8,2,4.1,5.9,5.9] if tajik else [3.0]*len(headers))
    for index,head in enumerate(headers):
        cell=table.rows[0].cells[index]
        cell.text=str(head)
        cell.width=Cm(widths[index])
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold=True
                run.font.color.rgb=RGBColor(24,54,97)
                run.font.size=Pt(8 if tajik else 6.4)
    # repeat table header on every page
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    trpr=table.rows[0]._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:tblHeader'))
    for data in rows:
        cells=table.add_row().cells
        for n,val in enumerate(data):
            s=str(val if val is not None else '')
            # Word text isn't evaluated as Excel formulas.
            cells[n].text=s
            for p in cells[n].paragraphs:
                for run in p.runs:
                    run.font.size=Pt(8 if tajik else 6.4)
    return table


def _write_docx(target: Path, *, rows: list[list[object]], tajik: bool,
                no_number: bool=False, template: bool=False) -> None:
    doc=Document()
    _style_doc(doc,tajik)
    h=doc.add_heading('Рӯйхати хонандагон' if tajik else 'Реестр учеников',level=1)
    for r in h.runs:r.font.size=Pt(16)
    note=('Заполните имя, фамилию и класс. Номер по порядку не требуется; ID создаётся программой.'
          if tajik else 'ID записи можно оставлять пустым; новые ID формируются автоматически при импорте.')
    p=doc.add_paragraph(note)
    p.style='Normal'
    if tajik:
        headers=TAJIK_HEADERS[1:] if no_number else TAJIK_HEADERS
    else:
        headers=HEADERS[1:] if template else HEADERS
    _table(doc,headers,rows,tajik)
    target.parent.mkdir(parents=True,exist_ok=True)
    doc.save(target)


def save_word_template(target: Path, *, tajik: bool=False, no_number: bool=False) -> None:
    target=Path(target)
    if target.suffix.lower()=='.doc':
        with tempfile.TemporaryDirectory(prefix='schoolregister_word_') as tmp:
            docx=Path(tmp)/(target.stem+'.docx')
            _write_docx(docx,rows=[['']* (5 if no_number and tajik else 6 if tajik else 17) for _ in range(8)],
                        tajik=tajik,no_number=no_number,template=True)
            convert_legacy_word(docx,target)
    else:
        _write_docx(target,rows=[['']* (5 if no_number and tajik else 6 if tajik else 17) for _ in range(8)],
                    tajik=tajik,no_number=no_number,template=True)


def export_word(db, target: Path, *, tajik: bool=False, no_number: bool=False) -> None:
    target=Path(target)
    if target.suffix.lower() not in ('.docx','.doc'):
        raise ValueError('Экспорт Word только в .docx или .doc')
    sql='SELECT * FROM students WHERE deleted_at IS NULL ORDER BY surname COLLATE NOCASE,given_name COLLATE NOCASE'
    rows=[]
    for i,row in enumerate(db.con.execute(sql),1):
        if tajik:
            name=' '.join(part for part in [row['surname'],row['given_name'],row['patronymic']] if part)
            vals=[name,str(row['grade'])+(row['section'] or ''),row['guardian_phone'],row['to_school'],row['from_school']]
            if not no_number:vals.insert(0,i)
        else:
            vals=[row[key] for key in FIELDS]
        rows.append(vals)
    if target.suffix.lower()=='.doc':
        with tempfile.TemporaryDirectory(prefix='schoolregister_word_') as tmp:
            docx=Path(tmp)/(target.stem+'.docx')
            _write_docx(docx,rows=rows,tajik=tajik,no_number=no_number)
            convert_legacy_word(docx,target)
    else:
        _write_docx(target,rows=rows,tajik=tajik,no_number=no_number)
